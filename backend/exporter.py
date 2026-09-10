import io
import datetime
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from .models import RecordResponse

def export_records_to_excel(records: List[RecordResponse]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "SAIL Material Records"

    # Title Styling
    title_font = Font(name="Calibri", size=14, bold=True, color="002855")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="555555")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="002855", end_color="002855", fill_type="solid")
    alt_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    # Title Rows
    ws.merge_cells("A1:O1")
    ws["A1"] = "STEEL AUTHORITY OF INDIA LIMITED - SALEM STEEL PLANT"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:O2")
    ws["A2"] = f"MATERIAL MANAGEMENT MODULE - DOCUMENT EXTRACTION REGISTER (Generated on {datetime.datetime.now().strftime('%d-%b-%Y %H:%M')})"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    # Column Headers
    headers = [
        "Record ID", "Document Type", "Document Ref ID", "Plant", "Supplier / Vendor",
        "Material Name", "Grade / Spec", "Quantity", "Unit", "Heat / Batch No",
        "PO Number", "Invoice Number", "Document Date", "Confidence (%)", "Remarks"
    ]

    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num)
        cell.value = header_title
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[4].height = 28

    # Data Rows
    for row_idx, r in enumerate(records, 5):
        row_values = [
            r.id,
            r.document_type,
            r.document_id,
            r.plant,
            r.supplier_name,
            r.material_name,
            r.material_grade_spec,
            r.quantity,
            r.unit,
            r.heat_batch_number,
            r.po_number,
            r.invoice_number,
            r.date_of_document,
            f"{r.confidence_score:.1f}%",
            r.remarks
        ]

        is_alt = (row_idx % 2 == 0)
        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = val
            cell.font = Font(name="Calibri", size=9)
            cell.border = thin_border
            if is_alt:
                cell.fill = alt_fill
            
            # Alignments
            if col_idx in [1, 8, 9, 10, 11, 12, 13, 14]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto adjust column widths safely
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max([len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(4, ws.max_row + 1)] or [10])
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def export_single_record_pdf(record: RecordResponse) -> io.BytesIO:
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#002855'),
        alignment=1,
        fontName='Helvetica-Bold'
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#4A5568'),
        alignment=1,
        fontName='Helvetica'
    )
    section_style = ParagraphStyle(
        'SectionHeader',
        parent=styles['Heading2'],
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#002855'),
        fontName='Helvetica-Bold',
        spaceBefore=10,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#2D3748')
    )

    story = []

    # Header
    story.append(Paragraph("STEEL AUTHORITY OF INDIA LIMITED", title_style))
    story.append(Paragraph("Salem Steel Plant ? Material Management Module", subtitle_style))
    story.append(Paragraph("Standardized Material Document Extraction Report", subtitle_style))
    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#002855'), spaceBefore=2, spaceAfter=12))

    # General Information Table
    general_data = [
        [
            Paragraph("<b>Document Ref ID:</b>", body_style), Paragraph(record.document_id, body_style),
            Paragraph("<b>Document Type:</b>", body_style), Paragraph(record.document_type, body_style)
        ],
        [
            Paragraph("<b>Steel Plant:</b>", body_style), Paragraph(record.plant, body_style),
            Paragraph("<b>Document Date:</b>", body_style), Paragraph(record.date_of_document, body_style)
        ],
        [
            Paragraph("<b>Original Filename:</b>", body_style), Paragraph(record.original_filename, body_style),
            Paragraph("<b>Extraction Timestamp:</b>", body_style), Paragraph(record.extracted_on[:19].replace('T', ' '), body_style)
        ]
    ]

    gen_table = Table(general_data, colWidths=[110, 150, 110, 150])
    gen_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(Paragraph("1. DOCUMENT METADATA", section_style))
    story.append(gen_table)
    story.append(Spacer(1, 12))

    # Material Specifications & Commercial Details
    mat_data = [
        [Paragraph("<b>Field Name</b>", body_style), Paragraph("<b>Extracted Value</b>", body_style), Paragraph("<b>Field Confidence</b>", body_style)],
        [Paragraph("<b>Supplier / Vendor</b>", body_style), Paragraph(record.supplier_name or "-", body_style), Paragraph(f"{record.field_confidence.get('supplier_name', 0):.0f}%", body_style)],
        [Paragraph("<b>Material Description</b>", body_style), Paragraph(record.material_name or "-", body_style), Paragraph(f"{record.field_confidence.get('material_name', 0):.0f}%", body_style)],
        [Paragraph("<b>Grade / Specification</b>", body_style), Paragraph(record.material_grade_spec or "-", body_style), Paragraph(f"{record.field_confidence.get('material_grade_spec', 0):.0f}%", body_style)],
        [Paragraph("<b>Quantity & Unit</b>", body_style), Paragraph(f"{record.quantity} {record.unit}", body_style), Paragraph(f"{record.field_confidence.get('quantity', 0):.0f}%", body_style)],
        [Paragraph("<b>Heat / Batch / Cast No</b>", body_style), Paragraph(record.heat_batch_number or "-", body_style), Paragraph(f"{record.field_confidence.get('heat_batch_number', 0):.0f}%", body_style)],
        [Paragraph("<b>Purchase Order No</b>", body_style), Paragraph(record.po_number or "-", body_style), Paragraph(f"{record.field_confidence.get('po_number', 0):.0f}%", body_style)],
        [Paragraph("<b>Invoice / Bill No</b>", body_style), Paragraph(record.invoice_number or "-", body_style), Paragraph(f"{record.field_confidence.get('invoice_number', 0):.0f}%", body_style)],
        [Paragraph("<b>Remarks / QC Status</b>", body_style), Paragraph(record.remarks or "-", body_style), Paragraph(f"{record.field_confidence.get('remarks', 0):.0f}%", body_style)],
    ]

    mat_table = Table(mat_data, colWidths=[150, 270, 100])
    mat_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#002855')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
    ]))
    story.append(Paragraph("2. MATERIAL & QUALITY CONTROL ATTRIBUTES", section_style))
    story.append(mat_table)
    story.append(Spacer(1, 14))

    conf_color = '#10B981' if record.confidence_score >= 85 else ('#F59E0B' if record.confidence_score >= 65 else '#EF4444')
    story.append(Paragraph(f"<b>Overall Extraction Confidence Score:</b> <font color='{conf_color}'><b>{record.confidence_score:.1f}%</b></font>", body_style))
    story.append(Spacer(1, 16))

    sign_data = [
        [
            Paragraph("<b>Extracted By:</b><br/>SAIL OCR Module v1.0", body_style),
            Paragraph("<b>Reviewed & Verified By:</b><br/>___________________________", body_style),
            Paragraph("<b>Store Officer / QC In-Charge:</b><br/>Salem Steel Plant", body_style)
        ]
    ]
    sign_table = Table(sign_data, colWidths=[170, 180, 170])
    sign_table.setStyle(TableStyle([
        ('TOPPADDING', (0, 0), (-1, -1), 10),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
        ('LINEABOVE', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1'))
    ]))
    story.append(sign_table)

    doc.build(story)
    output.seek(0)
    return output
