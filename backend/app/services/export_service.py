"""
Export Service for SAIL Material Management Module
Generates:
1. Excel (.xlsx) with fixed 9-section template + materials sheet
2. Official 2-Page SAIL Procurement Template PDF (.pdf) matching reference design
3. JSON (.json) matching Section 10 schema
"""
import io
import json
import pathlib
from typing import Dict, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.services.template_pdf_generator import generate_procurement_template_pdf
from app.services.template_docx_generator import generate_procurement_template_docx

class ExportService:
    def __init__(self):
        pass

    def _ensure_dict(self, data: Any) -> Dict[str, Any]:
        if isinstance(data, str):
            try:
                return json.loads(data)
            except Exception:
                return {}
        return data or {}

    def export_docx(self, structured_data: Any, filename: str) -> io.BytesIO:
        """Generates the official 2-page SAIL Salem Steel Plant Procurement Template editable DOCX."""
        data = self._ensure_dict(structured_data)
        return generate_procurement_template_docx(data, filename)

    def export_excel(self, structured_data: Any, filename: str) -> io.BytesIO:
        structured_data = self._ensure_dict(structured_data)
        wb = openpyxl.Workbook()
        ws_summary = wb.active
        ws_summary.title = "Procurement Summary"

        title_font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
        title_fill = PatternFill(start_color="101B24", end_color="101B24", fill_type="solid")
        
        section_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        section_fill = PatternFill(start_color="24313C", end_color="24313C", fill_type="solid")
        
        label_font = Font(name="Calibri", size=10, bold=True, color="334155")
        label_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        val_font = Font(name="Calibri", size=10, color="0F172A")

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        # Title Header
        ws_summary.merge_cells("A1:D1")
        ws_summary["A1"] = "SAIL MATERIAL MANAGEMENT MODULE — SALEM STEEL PLANT"
        ws_summary["A1"].font = title_font
        ws_summary["A1"].fill = title_fill
        ws_summary["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws_summary.row_dimensions[1].height = 30

        ws_summary.merge_cells("A2:D2")
        ws_summary["A2"] = f"Document Analysis Report: {filename}"
        ws_summary["A2"].font = Font(name="Calibri", size=11, italic=True, color="64748B")
        ws_summary["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws_summary.row_dimensions[2].height = 20

        current_row = 4

        sections = [
            ("1. Document Information", structured_data.get("document_information", {})),
            ("2. Material Information", structured_data.get("material_information", {})),
            ("3. Quantity Information", structured_data.get("quantity_information", {})),
            ("4. Procurement Information", structured_data.get("procurement_information", {})),
            ("5. Technical Information", structured_data.get("technical_information", {})),
            ("6. Commercial Information", structured_data.get("commercial_information", {})),
            ("7. Additional Information", structured_data.get("additional_information", {})),
            ("8. OCR / AI Confidence", structured_data.get("confidence", {})),
            ("9. Source Document", structured_data.get("source_document", {}))
        ]

        for sec_title, sec_data in sections:
            ws_summary.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
            cell = ws_summary.cell(row=current_row, column=1, value=sec_title)
            cell.font = section_font
            cell.fill = section_fill
            cell.alignment = Alignment(vertical="center", indent=1)
            ws_summary.row_dimensions[current_row].height = 22
            current_row += 1

            for k, v in sec_data.items():
                label_text = k.replace("_", " ").title()
                val_text = str(v)

                lbl_cell = ws_summary.cell(row=current_row, column=1, value=label_text)
                lbl_cell.font = label_font
                lbl_cell.fill = label_fill
                lbl_cell.border = thin_border

                ws_summary.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=4)
                val_cell = ws_summary.cell(row=current_row, column=2, value=val_text)
                val_cell.font = val_font
                val_cell.border = thin_border
                val_cell.alignment = Alignment(wrap_text=True)
                ws_summary.row_dimensions[current_row].height = 20
                current_row += 1

            current_row += 1

        ws_summary.column_dimensions["A"].width = 30
        ws_summary.column_dimensions["B"].width = 30
        ws_summary.column_dimensions["C"].width = 30
        ws_summary.column_dimensions["D"].width = 30

        # Sheet 2: Extracted Materials Table
        ws_mat = wb.create_sheet(title="Extracted Materials")
        ws_mat.merge_cells("A1:I1")
        ws_mat["A1"] = "EXTRACTED MATERIALS TABLE — SALEM STEEL PLANT"
        ws_mat["A1"].font = title_font
        ws_mat["A1"].fill = title_fill
        ws_mat["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws_mat.row_dimensions[1].height = 28

        mat_headers = ["Sl.No", "Material Code", "Material Description", "Specification", "Quantity", "Unit", "Grade", "Make / Brand", "Remarks"]
        for col_idx, h in enumerate(mat_headers, 1):
            c = ws_mat.cell(row=3, column=col_idx, value=h)
            c.font = section_font
            c.fill = section_fill
            c.border = thin_border
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws_mat.row_dimensions[3].height = 24

        materials = structured_data.get("materials", [])
        m_row = 4
        for m in materials:
            vals = [
                m.get("sl_no", m_row - 3),
                m.get("material_code", "Not Available"),
                m.get("material_description", "Not Available"),
                m.get("specification", "Not Available"),
                m.get("quantity", "Not Available"),
                m.get("unit", "Not Available"),
                m.get("grade", "Not Available"),
                m.get("make_brand", "Not Available"),
                m.get("remarks", "Not Available")
            ]
            for col_idx, val in enumerate(vals, 1):
                c = ws_mat.cell(row=m_row, column=col_idx, value=str(val))
                c.font = val_font
                c.border = thin_border
                c.alignment = Alignment(vertical="center")
            ws_mat.row_dimensions[m_row].height = 20
            m_row += 1

        for col in ["A", "B", "C", "D", "E", "F", "G", "H", "I"]:
            ws_mat.column_dimensions[col].width = 20
        ws_mat.column_dimensions["C"].width = 35
        ws_mat.column_dimensions["D"].width = 35

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def export_pdf(self, structured_data: Any, filename: str) -> io.BytesIO:
        """Generates the official 2-page SAIL Salem Steel Plant Procurement Template PDF."""
        data = self._ensure_dict(structured_data)
        return generate_procurement_template_pdf(data, filename)

    def export_json(self, structured_data: Any) -> str:
        data = self._ensure_dict(structured_data)
        return json.dumps(data, indent=2)

export_service = ExportService()
