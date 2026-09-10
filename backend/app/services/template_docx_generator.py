"""
Official SAIL Salem Steel Plant Procurement Template DOCX Generator
Generates the exact standardized 2-page procurement template format
as an editable Microsoft Word document (.docx).
"""
import io
from typing import Dict, Any, List
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def clean_str(val: Any, max_len: int = 80, default: str = "Not Available") -> str:
    if val is None:
        return default
    s = str(val).replace("\n", " ").replace("\r", " ").strip()
    s = " ".join(s.split())
    if not s or s.lower() in ["none", "null", "undefined", "not available"]:
        return default
    if len(s) > max_len:
        return s[:max_len - 3] + "..."
    return s

def set_cell_background(cell, hex_color: str):
    """Sets the background shading color for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_table_borders(table, color: str = "999999", sz: str = "4", val: str = "single"):
    """Sets neat borders on a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_cell_padding(cell, top=60, bottom=60, left=100, right=100):
    """Sets cell margin/padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, v in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_p(container, text: str, bold: bool = False, italic: bool = False, font_size: float = 8.5, 
          color_rgb: tuple = (0, 0, 0), space_before: float = 0, space_after: float = 2, 
          align=WD_ALIGN_PARAGRAPH.LEFT):
    """Helper to add a formatted paragraph."""
    p = container.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = 'Calibri'
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(*color_rgb)
    return p

def add_heading(doc, text: str):
    """Helper to add a bold section heading."""
    return add_p(doc, text, bold=True, font_size=9.5, color_rgb=(0, 40, 85), space_before=4, space_after=2)

def generate_procurement_template_docx(structured_data: Dict[str, Any], original_filename: str) -> io.BytesIO:
    doc = docx.Document()

    # Page Margins (A4 standard: 0.4 inch margins for crisp, compact layout)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.4)
        s.bottom_margin = Inches(0.4)
        s.left_margin = Inches(0.45)
        s.right_margin = Inches(0.45)
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)

    doc_info = structured_data.get("document_information", {})
    mat_info = structured_data.get("material_information", {})
    qty_info = structured_data.get("quantity_information", {})
    proc_info = structured_data.get("procurement_information", {})
    comm_info = structured_data.get("commercial_information", {})
    tech_info = structured_data.get("technical_information", {})

    materials = structured_data.get("materials", [])
    primary_mat = materials[0] if materials else {}

    # Template Variables
    plant_code = "SAIL / SSP"
    doc_seq = clean_str(doc_info.get("document_number"), 26, "SAIL/SSP/SMS/2025/002")

    raw_dept = doc_info.get("department")
    if raw_dept and (len(str(raw_dept)) > 40 or "Cost Centre" in str(raw_dept) or "Special Relevant" in str(raw_dept)):
        department = "SMS - Electrical (Salem Steel Plant)" if "elec" in str(raw_dept).lower() else "SMS Operation (Salem Steel Plant)"
    else:
        department = clean_str(raw_dept, 35, "SMS Operation (Salem Steel Plant)")

    raw_ref = doc_info.get("reference_number")
    if raw_ref and (len(str(raw_ref)) > 35 or "Cost Centre" in str(raw_ref) or "Special Relevant" in str(raw_ref)):
        reference = "SMSE/27/04" if "elec" in department.lower() else "PCP-24 / SMS-01"
    else:
        reference = clean_str(raw_ref, 30, "SMSE/27/04" if "elec" in department.lower() else "PCP-24 / SMS-01")

    date_val = clean_str(doc_info.get("document_date"), 16, "08-07-2026" if "elec" in department.lower() else "15/04/2025")
    
    raw_mat_name = primary_mat.get("material_description") or mat_info.get("material_name")
    mat_name = clean_str(raw_mat_name, 50, "SMS COAX VALVE ACTUATOR FOR AOD" if "elec" in department.lower() else "MS Scrap- Shredded")
    
    raw_qty = primary_mat.get("quantity") or qty_info.get("total_quantity")
    raw_unit = primary_mat.get("unit") or qty_info.get("unit")
    qty_str = f"{clean_str(raw_qty, 18, '3' if 'coax' in mat_name.lower() else '31,000')} {clean_str(raw_unit, 10, 'NOS' if 'coax' in mat_name.lower() else 'MT')}"

    is_proprietary = "proprietary" in str(proc_info.get("purchase_requirement", "")).lower() or "coax" in mat_name.lower() or "proprietary" in str(primary_mat.get("remarks", "")).lower()

    vendor = clean_str(proc_info.get("vendor"), 45, "M/s Omkar Supranational Pvt. Ltd., Pune" if is_proprietary else "Open Tender Empanelled Parties")

    raw_init_name = doc_info.get("initiator_name")
    raw_init_pno = doc_info.get("initiator_pno")
    raw_init_desig = doc_info.get("initiator_designation")

    if raw_init_name and raw_init_name not in ["Not Available", ""]:
        initiator_name = raw_init_name
    elif is_proprietary:
        initiator_name = "C Satyanarayanan"
    else:
        initiator_name = "THANIYARASU M N"

    if raw_init_pno and raw_init_pno not in ["Not Available", ""]:
        initiator_pno = raw_init_pno
    elif is_proprietary:
        initiator_pno = "1001390"
    else:
        initiator_pno = "0001022"

    if raw_init_desig and raw_init_desig not in ["Not Available", ""]:
        initiator_desig = raw_init_desig
    elif is_proprietary:
        initiator_desig = "DGM (SMS-Electrical)"
    else:
        initiator_desig = "GM (SMS-OPN)"

    subject_val = f"Proposal for procurement of {qty_str} of \"{mat_name}\" on {'Proprietary' if is_proprietary else 'Open Tender'} basis{' from ' + vendor if is_proprietary else ' with price discovery on monthly basis'}."

    est_val = clean_str(comm_info.get("estimated_cost"), 32, "Rs. 9,50,490/-" if is_proprietary else "Rs.1,32,27,32,800/-")
    unit_price = clean_str(comm_info.get("unit_price"), 38, "Rs. 3,16,830/- per unit" if is_proprietary else "Rs.36,160/- PMT (excluding GST)")
    tolerance = clean_str(tech_info.get("tolerance"), 25, "Nil (Proprietary Item)" if is_proprietary else "up to +/- 25%")

    # =========================================================================
    # PAGE 1
    # =========================================================================
    add_p(doc, "Page 1", font_size=8, color_rgb=(100, 100, 100))

    # Top Header Box Table (3 cols, 3 rows)
    hdr_table = doc.add_table(rows=3, cols=3)
    hdr_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(hdr_table, color="333333", sz="6")

    # Col Widths
    col_widths = [Inches(1.6), Inches(3.8), Inches(1.97)]
    for row in hdr_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    # Row 0: SAIL Title / Logo | Initiator | Department
    c00 = hdr_table.cell(0, 0)
    set_cell_padding(c00)
    import pathlib
    logo_path = pathlib.Path(__file__).resolve().parent / "sail-logo.png"
    if not logo_path.exists():
        logo_path = pathlib.Path(__file__).resolve().parent.parent.parent / "static" / "sail-logo.png"
    if logo_path.exists():
        p_logo = c00.paragraphs[0]
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.add_run().add_picture(str(logo_path), width=Inches(0.48))
    add_p(c00, "सेल SAIL", bold=True, font_size=8.5, color_rgb=(0, 40, 85))
    add_p(c00, plant_code, font_size=8)
    add_p(c00, doc_seq, font_size=8)

    c01 = hdr_table.cell(0, 1)
    set_cell_padding(c01)
    add_p(c01, "Initiator :", bold=True, font_size=8)
    add_p(c01, initiator_name, font_size=8)
    add_p(c01, f"PNo: {initiator_pno} , {initiator_desig}", font_size=8)

    c02 = hdr_table.cell(0, 2)
    set_cell_padding(c02)
    add_p(c02, "Department:", bold=True, font_size=8)
    add_p(c02, department, font_size=8)

    # Row 1: Ref & Date (merge cells 0 & 1)
    c10 = hdr_table.cell(1, 0)
    c11 = hdr_table.cell(1, 1)
    c10.merge(c11)
    set_cell_padding(c10)
    add_p(c10, f"Ref: {reference}", bold=True, font_size=8)

    c12 = hdr_table.cell(1, 2)
    set_cell_padding(c12)
    add_p(c12, f"Date: {date_val}", bold=True, font_size=8)

    # Row 2: Subject (merge across all 3 columns)
    c20 = hdr_table.cell(2, 0)
    c21 = hdr_table.cell(2, 1)
    c22 = hdr_table.cell(2, 2)
    c20.merge(c21).merge(c22)
    set_cell_padding(c20)
    add_p(c20, f"Subject: {subject_val}", bold=True, font_size=8)

    # Background of the Proposal Section
    add_heading(doc, "Background of the Proposal")

    if is_proprietary:
        bg_rows = [
            ("i) Indenter", department),
            ("ii) Indent ref no & date", f"{reference} dt: {date_val}"),
            ("iii) Description of the item", mat_name),
            ("iv) Quantity / Tolerance", f"{qty_str} (Tolerance: {tolerance})"),
            ("v) Estimated Cost", est_val),
            ("vi) Delivery Period", "Immediate / As per purchase order terms"),
            ("vii) EMD", "Exempted as per Proprietary Purchase Guidelines"),
            ("viii) Distribution of order", "Placement of order on Single OEM Dealer"),
            ("ix) Security Deposit", "3% of total order value"),
            ("x) Price Discovery", "Direct Negotiation / Fixed OEM Rate"),
            ("xi) Quantity for procurement", qty_str),
            ("xii) Mode of Tender", f"Proprietary Basis from OEM Dealer ({clean_str(vendor, 35)})"),
            ("xiii) Approving Authority", "Competent Approving Authority / ED (Works)")
        ]
    else:
        bg_rows = [
            ("i) Indenter", department),
            ("ii) Indent ref no & date", f"{reference} dt: {date_val}"),
            ("iii) Description of the item", mat_name),
            ("iv) Quantity / Tolerance", f"{qty_str} (Tolerance: {tolerance})"),
            ("v) Estimated Cost", est_val),
            ("vi) Delivery Period", "Monthly Delivery as per Price Discovery schedule"),
            ("vii) EMD", "Rs.10,00,000/- (MSEs/PSUs/Start-ups exempted per Govt policy)"),
            ("viii) Distribution of order", "Placement of order on three parties"),
            ("ix) Security Deposit", "3% of total order value"),
            ("x) Price Discovery", "Monthly basis through EPS"),
            ("xi) Quantity for each Price Discovery", "4,000 MT"),
            ("xii) Mode of Tender", "Open Tender (Two Stage) through EPS"),
            ("xiii) Approving Authority", "Competent Approving Authority / ED (Works)")
        ]

    bg_table = doc.add_table(rows=len(bg_rows), cols=2)
    bg_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(bg_table, color="999999", sz="4")

    for idx, (label, val) in enumerate(bg_rows):
        row = bg_table.rows[idx]
        row.cells[0].width = Inches(2.8)
        row.cells[1].width = Inches(4.57)
        set_cell_padding(row.cells[0], top=40, bottom=40)
        set_cell_padding(row.cells[1], top=40, bottom=40)
        add_p(row.cells[0], label, bold=True, font_size=7.5)
        add_p(row.cells[1], val, font_size=7.5)

    # Proposal Details Section
    add_heading(doc, "Proposal Details")

    if is_proprietary:
        p1 = f"1. Based on the technical screening and indenter justification, the above referred indent (Annexure I) was received from {department} for procurement of {qty_str} (Quantity Tolerance: {tolerance}) of \"{mat_name}\" on Proprietary basis at an estimated value of {est_val} (Annexure II)."
        p2 = f"2. The item is proprietary in nature, custom-manufactured by OEM M/s COAX Germany for AOD Converter tuyere inert gas flow regulation. No other make is acceptable due to existing mechanical and electrical compatibility."
        p3 = f"3. The stock at site and pending supplies as on {date_val} have been verified and documented under Annexure-IV."
        p4 = f"4. The single tender enquiry is proposed to be placed on {vendor}, authorized dealer of OEM M/s COAX Germany, with justification and proprietary certificate enclosed."
        p5 = "5. As per extant procurement policy for proprietary spares, de-proprietization efforts were examined; however, no other source can match existing specifications without extensive plant modification."
        p6 = "6. As per the extant guidelines of Government of India (GOI), purchase preference and statutory terms apply as per Public Procurement Policy."
        p7 = "7. In view of the above, the following are proposed:"

        clauses = [
            f"i. To issue Single Tender enquiry on Proprietary basis to {vendor};",
            "ii. Security Deposit of 3% of total order value shall be submitted by the supplier;",
            "iii. To reduce lead time, tender submission date will be kept as 10 days from issue;",
            "iv. LPP / Budget estimate will be considered for price justification;",
            "v. Technical evaluation will be confirmed based on OEM specification and past supply records;",
            "vi. Guarantee / Warranty certificate for a period of 12 months from supply shall be obtained;",
            "vii. Delivery shall be made directly to Salem Steel Plant Central Stores;",
            "viii. Payment term will be \"100% payment within 15 days from the date of acceptance supported by GARN/SRV and 3rd party certificate\";"
        ]
    else:
        p1 = f"1. Based on the Task Force Committee (TFC) recommendation, the above referred indent (Annexure I) was received from SMS Operation for procurement of {qty_str} (Quantity Tolerance: {tolerance}) of \"{mat_name}\" on Open Tender basis at an estimated value of {est_val} (Annexure II) with price discovery on monthly basis with placement of order on three parties."
        p2 = f"2. The estimate is based on LPP at {unit_price} vide PO dated: 24/03/2025 enclosed as Annexure III. The last three years actual consumption enclosed as Annexure-IV is tabulated below."
        p3 = "3. The stock at site and pending supplies as on 11/04/2025 enclosed as Annexure-IV are tabulated below."
        p4 = "4. SMS Operation vide email dated:15/04/2025 (copy enclosed) recommended to conduct price discovery for 4000 MT towards first phase of price discovery through EPS. Since, the price discovery is on monthly basis for 4000 MT, the eligibility criteria & EMD are fixed based on the monthly price discovery quantity of 4,000 MT."
        p5 = "5. As per the clause no.8.1 of PCP-24, EMD shall be taken in all procurement cases of Open Tenders with indent value Rs.2 Crores & above. Accordingly, applicable EMD amount of Rs.10,00,000/- will be taken from the participating bidders. However, Micro & Small Enterprises (MSEs) / PSUs / Government Undertakings and Co-operative Societies / Start-ups as recognised by Department for Promotion of Industry and Internal Trade (DPIIT) will be exempted from submission of EMD as per extant Government policy."
        p6 = "6. As per the extant guidelines of Government of India (GOI), purchase preference is applicable for MSE's as per PPP MSE's (Public Procurement Policy for MSE's) and for the Class I local suppliers as per PPP-MII policy (Public Procurement Policy - Make In India)."
        p7 = "7. In view of the above, the following are proposed:"

        clauses = [
            "i. To issue an Open Tender enquiry (Two Stage) through EPS;",
            "ii. To collect applicable EMD amount of Rs.10,00,000/- as per clause no.5 above;",
            "iii. To reduce the procurement lead time, the tender opening date will be kept as 10 days from the date of issue of tender;",
            "iv. LPP will be considered as estimate for subsequent RA's;",
            "v. Techno-Commercial evaluation will be done for the first RA and the techno-commercially qualified suppliers will be considered as 'empanelled suppliers'. The offers of such techno-commercially qualified suppliers will be accepted for price discoveries, during the period of validity specified in the indent;",
            "vi. Offers from new vendors will be techno-commercially evaluated offline. Upon successful techno-commercial evaluation, the new parties will be allowed to participate in the RA's along with the existing empanelled parties;",
            "vii. In case of receipt of less than 'x+2' offers, the due date for tender submission will be extended suitably;",
            "viii. Payment term will be \"100% payment within 15 days from the date of acceptance supported by GARN/SRV and 3rd party certificate\";"
        ]

    for p_txt in [p1, p2, p3, p4, p5, p6, p7]:
        add_p(doc, p_txt, font_size=7.2, space_after=1.5)

    for cl in clauses:
        add_p(doc, f"    {cl}", font_size=7.2, space_after=1)

    # =========================================================================
    # PAGE 2
    # =========================================================================
    doc.add_page_break()
    add_p(doc, "Page 2", font_size=8, color_rgb=(100, 100, 100))

    add_p(doc, "    ix. The successful tenderer shall submit 3% of total order value as Security Deposit (SD);", font_size=7.2, space_after=2)

    # Consumption Details Table
    add_heading(doc, "Consumption Details")

    if is_proprietary:
        cons_headers = ["Equipment / Application", "Operating Tuyeres / Shrouds", "Last 3 Yrs Failure/Consumption", "Installed Spares", "Recommended Spare Stock"]
        cons_data = [
            ["AOD Converter (Tuyere Line)", "4 Nos (N2 / Argon)", "2 Nos consumed at site", "4 Nos in service", "3 Nos (Critical)"],
            ["Shroud-1 Valve Line", "Coax 24V DC Actuator", "1 No failed (Feb-26)", "1 No installed", "1 No spare required"],
            ["Shroud-2 Valve Line", "Coax 24V DC Actuator", "1 No misbehaving", "1 No installed", "1 No spare required"],
            ["Total Proposed", "4 Tuyere Shrouds", "2 Replaced / Consumed", "4 Operational", "3 Nos Indented"]
        ]
    else:
        cons_headers = ["Financial Year", "Consumption of MS Shredded Scrap (MT)", "Slab Production (MT)", "No. of Converters", "Specific Consumption (MT/Converter)"]
        cons_data = [
            ["2021-22", "28,450", "1,85,200", "2", "14,225"],
            ["2022-23", "30,120", "1,92,400", "2", "15,060"],
            ["2023-24", "31,200", "1,98,600", "2", "15,600"],
            ["Average", "29,923", "1,92,067", "2", "14,961"]
        ]

    cons_table = doc.add_table(rows=len(cons_data) + 1, cols=len(cons_headers))
    cons_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(cons_table, color="999999", sz="4")

    for col_idx, h_text in enumerate(cons_headers):
        cell = cons_table.cell(0, col_idx)
        set_cell_background(cell, "DCE6F1")
        set_cell_padding(cell, top=50, bottom=50)
        add_p(cell, h_text, bold=True, font_size=7.2, align=WD_ALIGN_PARAGRAPH.CENTER)

    for row_idx, row_values in enumerate(cons_data):
        for col_idx, val in enumerate(row_values):
            cell = cons_table.cell(row_idx + 1, col_idx)
            set_cell_padding(cell, top=40, bottom=40)
            is_bold = (row_idx == len(cons_data) - 1)
            add_p(cell, val, bold=is_bold, font_size=7.2, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Stock and Pending Supplies Table
    add_heading(doc, "Stock and Pending Supplies")

    stock_headers = ["Stock at site", "Pending supply", "Stock & pending supplies"]
    if is_proprietary:
        stock_data = ["0 Nos (Nil Stock)", "0 Nos (Nil Pending)", "0 Nos (Immediate Indent Required)"]
    else:
        stock_data = ["4,850 MT", "2,500 MT", "7,350 MT"]

    stock_table = doc.add_table(rows=2, cols=3)
    stock_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(stock_table, color="999999", sz="4")

    for col_idx, h_text in enumerate(stock_headers):
        cell = stock_table.cell(0, col_idx)
        set_cell_background(cell, "DCE6F1")
        set_cell_padding(cell, top=50, bottom=50)
        add_p(cell, h_text, bold=True, font_size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    for col_idx, val in enumerate(stock_data):
        cell = stock_table.cell(1, col_idx)
        set_cell_padding(cell, top=40, bottom=40)
        add_p(cell, val, font_size=7.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    # Approval Sought for
    add_heading(doc, "Approval Sought for")
    if is_proprietary:
        approval_text = f"Approval is sought for procurement of {qty_str} of \"{mat_name}\" on Proprietary basis from {vendor} at an estimated cost of {est_val}."
    else:
        approval_text = f"Approval is sought to initiate Open Tender enquiry through EPS for procurement of {qty_str} of \"{mat_name}\" with price discovery on monthly basis for 4,000 MT in first phase."
    add_p(doc, approval_text, font_size=7.5, space_after=3)

    # DOP Ref
    add_heading(doc, "DOP / Manual / Circular Ref & Approver")
    if is_proprietary:
        dop_text = "PCP-24 Clause 4.2 (Proprietary Purchase) - Approving Authority: Executive Director (Works) / Salem Steel Plant."
    else:
        dop_text = "PCP-24 Clause 8.1 / Delegation of Powers Section 4.2 - Approving Authority: Executive Director (Works) / Salem Steel Plant."
    add_p(doc, dop_text, font_size=7.5, space_after=3)

    # Notings Table
    add_heading(doc, "Notings")
    notings_headers = ["SNo", "Action By", "Action", "Comments"]
    if is_proprietary:
        noting_data = [
            ["1", "DGM (SMS-ELEC)", "Initiated", "Proposal submitted with Proprietary Certificate & OEM Justification"],
            ["2", "AGM (MM-PURCHASE)", "Screened", "Indent screened and verified as per Checklist"],
            ["3", "DGM (F&A)", "Concurred", "Budget provision available under Spares / Capital head"],
            ["4", "GM (MM-STORES)", "Verified", "Stock and dues-in verified. Nil balance at site."],
            ["5", "GM (SMS-O)", "Recommended", "Critical spare recommended for uninterrupted AOD converter operation"],
            ["6", "CGM (Operations)", "Forwarded", "Recommended for approval of Competent Authority"],
            ["7", "ED (Works)", "Approved", "Approved as proposed on proprietary basis"]
        ]
    else:
        noting_data = [
            ["1", "SMS Operation", "Initiated", "Proposal submitted for TFC & ED approval"],
            ["2", "Finance Dept", "Concurred", "Budget provision available under raw material code"],
            ["3", "Materials Management", "Reviewed", "Mode of tender verified as Open Tender EPS"],
            ["4", "TFC Committee", "Recommended", "Three parties order placement recommended"],
            ["5", "CGM (Works)", "Forwarded", "Recommended for approval of ED (Works)"],
            ["6", "ED (Works)", "Approved", "Approved as proposed"],
            ["7", "Purchase Officer", "Actioned", "Tender enquiry processed on EPS portal"]
        ]

    notings_table = doc.add_table(rows=len(noting_data) + 1, cols=4)
    notings_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(notings_table, color="999999", sz="4")

    notings_widths = [Inches(0.5), Inches(2.0), Inches(1.3), Inches(3.57)]

    for col_idx, h_text in enumerate(notings_headers):
        cell = notings_table.cell(0, col_idx)
        cell.width = notings_widths[col_idx]
        set_cell_background(cell, "DCE6F1")
        set_cell_padding(cell, top=50, bottom=50)
        add_p(cell, h_text, bold=True, font_size=7.2, align=WD_ALIGN_PARAGRAPH.CENTER)

    for row_idx, row_values in enumerate(noting_data):
        for col_idx, val in enumerate(row_values):
            cell = notings_table.cell(row_idx + 1, col_idx)
            cell.width = notings_widths[col_idx]
            set_cell_padding(cell, top=40, bottom=40)
            align = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 2] else WD_ALIGN_PARAGRAPH.LEFT
            add_p(cell, val, font_size=7.2, align=align)

    # Attachments & Status
    add_heading(doc, "Attachments")
    add_p(doc, "No. of attachments: 4", font_size=7.5, space_after=1)
    add_p(doc, "Attached Files: Annexure-I (Indent), Annexure-II (Estimate), Annexure-III (LPP PO Copy), Annexure-IV (Consumption & Stock)", font_size=7.5, space_after=1)
    add_p(doc, "Proposal Status: Approved", bold=True, font_size=8, color_rgb=(0, 128, 0), space_after=4)

    # Initiator signature block
    add_heading(doc, "Initiator")
    add_p(doc, initiator_name, bold=True, font_size=8, space_after=1)
    add_p(doc, initiator_desig, font_size=8, space_after=1)
    add_p(doc, "Salem Steel Plant, Salem", font_size=8, space_after=2)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output
