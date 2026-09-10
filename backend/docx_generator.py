import io
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from .models import EnquiryProposalNote

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def generate_proposal_docx(note: EnquiryProposalNote) -> io.BytesIO:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

    navy = RGBColor(0, 40, 85)
    dark_gray = RGBColor(51, 65, 85)

    # Top Header
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_ref = p_top.add_run(f"SAIL — Salem Steel Plant     {note.ref_no}\n")
    r_ref.font.size = Pt(8.5)
    r_ref.font.color.rgb = RGBColor(100, 116, 139)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t1 = p_title.add_run("SAIL — Salem Steel Plant\n")
    r_t1.font.size = Pt(13.5)
    r_t1.font.bold = True
    r_t1.font.color.rgb = navy

    r_t2 = p_title.add_run("ENQUIRY PROPOSAL NOTE\n")
    r_t2.font.size = Pt(11)
    r_t2.font.bold = True
    r_t2.font.color.rgb = navy

    # Metadata
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(8)
    p_meta.paragraph_format.line_spacing = 1.15

    def add_meta(lbl, val):
        r1 = p_meta.add_run(f"{lbl}: ")
        r1.bold = True
        r1.font.size = Pt(9)
        r1.font.color.rgb = navy
        r2 = p_meta.add_run(f"{val}\n")
        r2.font.size = Pt(9)
        r2.font.color.rgb = dark_gray

    add_meta("Initiator", note.initiator)
    add_meta("Department", note.department)
    add_meta("Ref", note.ref_no)
    add_meta("Date", note.date_of_document)
    add_meta("Subject", note.subject)

    # 13-Point Background Table
    p_bg = doc.add_paragraph()
    p_bg.paragraph_format.space_before = Pt(6)
    p_bg.paragraph_format.space_after = Pt(4)
    r_bg = p_bg.add_run("Background of the Proposal                                                                 ")
    r_bg.bold = True
    r_bg.font.size = Pt(10)
    r_bg.font.color.rgb = navy
    r_tag = p_bg.add_run("[EXTRACTED FROM YOUR FILE]")
    r_tag.font.size = Pt(7.5)
    r_tag.font.color.rgb = RGBColor(180, 83, 9)

    bg_items = [
        ("i)", "Indenter", note.indenter),
        ("ii)", "Indent Ref No & Date", note.indent_ref_date),
        ("iii)", "Description of the Item", note.item_description),
        ("iv)", "Material Code", note.material_code),
        ("v)", "Quantity", f"{note.quantity} {note.unit}"),
        ("vi)", "Estimated Cost", note.estimated_cost),
        ("vii)", "Delivery Period", note.delivery_period),
        ("viii)", "EMD", note.emd),
        ("ix)", "Distribution of Order", note.distribution_order),
        ("x)", "Security Deposit", note.security_deposit),
        ("xi)", "Price Discovery", note.price_discovery),
        ("xii)", "Mode of Tender", note.mode_of_tender),
        ("xiii)", "Approving Authority", note.approving_authority),
    ]

    t_bg = doc.add_table(rows=len(bg_items), cols=3)
    t_bg.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_bg)

    col_widths = [Inches(0.4), Inches(2.2), Inches(4.5)]
    for row_idx, (num, label, val) in enumerate(bg_items):
        row = t_bg.rows[row_idx]
        for c_idx, w in enumerate(col_widths):
            row.cells[c_idx].width = w
            set_cell_margins(row.cells[c_idx], top=50, bottom=50, left=80, right=80)
            if row_idx % 2 == 0:
                set_cell_background(row.cells[c_idx], "F8FAFC")

        row.cells[0].paragraphs[0].text = num
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(8)
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        row.cells[0].paragraphs[0].runs[0].font.color.rgb = navy

        row.cells[1].paragraphs[0].text = label
        row.cells[1].paragraphs[0].runs[0].font.size = Pt(8)
        row.cells[1].paragraphs[0].runs[0].font.bold = True

        row.cells[2].paragraphs[0].text = val
        row.cells[2].paragraphs[0].runs[0].font.size = Pt(8)

    # Section 1: Consumption Table
    p_sec1 = doc.add_paragraph()
    p_sec1.paragraph_format.space_before = Pt(10)
    p_sec1.paragraph_format.space_after = Pt(4)
    r1 = p_sec1.add_run("1. Last 3 years actual consumption, current stock, and pending supply details for MS-Shredded Scrap:\n")
    r1.bold = True
    r1.font.size = Pt(9)
    r1.font.color.rgb = navy
    r1_sub = p_sec1.add_run(note.consumption_intro)
    r1_sub.font.size = Pt(8)

    cons_hdr = ["Fin Year", "Consumption of MS-Shredded Scrap (MT)", "Slab Production (MT)", "No of Converters", "Specific Consumption (MT/ converter)"]
    t_cons = doc.add_table(rows=len(note.consumption_table) + 1, cols=5)
    t_cons.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_cons)

    for c_idx, h in enumerate(cons_hdr):
        cell = t_cons.rows[0].cells[c_idx]
        cell.text = h
        set_cell_background(cell, "002855")
        set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
        p = cell.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(7.5)

    for r_idx, crow in enumerate(note.consumption_table):
        row = t_cons.rows[r_idx + 1]
        vals = [crow.fin_year, crow.scrap_consumption, crow.slab_production, crow.no_converters, crow.specific_consumption]
        for c_idx, v in enumerate(vals):
            cell = row.cells[c_idx]
            cell.text = v
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            if r_idx % 2 == 1: set_cell_background(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(8)
            if crow.fin_year == "Average": p.runs[0].font.bold = True

    # Section 2: Requirement Table
    p_sec2 = doc.add_paragraph()
    p_sec2.paragraph_format.space_before = Pt(10)
    p_sec2.paragraph_format.space_after = Pt(4)
    r2 = p_sec2.add_run("2. Reason for deviation from consumption values:\n")
    r2.bold = True
    r2.font.size = Pt(9)
    r2.font.color.rgb = navy
    r2_sub = p_sec2.add_run(note.requirement_intro)
    r2_sub.font.size = Pt(8)

    req_hdr = ["Item", "Annual requirement (MT)", "Buffer stock (MT)", "Total (MT)"]
    t_req = doc.add_table(rows=len(note.requirement_table) + 1, cols=4)
    t_req.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_req)
    for c_idx, h in enumerate(req_hdr):
        cell = t_req.rows[0].cells[c_idx]
        cell.text = h
        set_cell_background(cell, "002855")
        set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
        p = cell.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(7.5)

    for r_idx, rrow in enumerate(note.requirement_table):
        row = t_req.rows[r_idx + 1]
        vals = [rrow.item, rrow.annual_req, rrow.buffer_stock, rrow.total]
        for c_idx, v in enumerate(vals):
            cell = row.cells[c_idx]
            cell.text = v
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            if r_idx % 2 == 1: set_cell_background(cell, "F8FAFC")
            row.cells[c_idx].paragraphs[0].runs[0].font.size = Pt(8)

    # Section 3: Justification Table
    p_sec3 = doc.add_paragraph()
    p_sec3.paragraph_format.space_before = Pt(10)
    p_sec3.paragraph_format.space_after = Pt(4)
    r3 = p_sec3.add_run(f"3. {note.justification_intro}")
    r3.bold = True
    r3.font.size = Pt(9)
    r3.font.color.rgb = navy

    just_hdr = ["Sl. No.", "Description", "Value"]
    t_just = doc.add_table(rows=len(note.justification_table) + 1, cols=3)
    t_just.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_just)
    for c_idx, h in enumerate(just_hdr):
        cell = t_just.rows[0].cells[c_idx]
        cell.text = h
        set_cell_background(cell, "002855")
        set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
        p = cell.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(7.5)

    for r_idx, jstep in enumerate(note.justification_table):
        row = t_just.rows[r_idx + 1]
        vals = [jstep.step_no, jstep.description, jstep.value]
        for c_idx, v in enumerate(vals):
            cell = row.cells[c_idx]
            cell.text = v
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            if r_idx % 2 == 1: set_cell_background(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(8)
            if jstep.step_no == "12": p.runs[0].font.bold = True

    # Section 4: Task Force Table
    p_sec4 = doc.add_paragraph()
    p_sec4.paragraph_format.space_before = Pt(10)
    p_sec4.paragraph_format.space_after = Pt(4)
    r4 = p_sec4.add_run(f"4. {note.task_force_intro}")
    r4.bold = True
    r4.font.size = Pt(9)
    r4.font.color.rgb = navy

    tf_hdr = ["Sl. No.", "Material code", "Item name", "Requirement as per ABP for FY 2025-26 plus Safety stock", "Total Quantity recommended for procurement in MT"]
    t_tf = doc.add_table(rows=len(note.task_force_table) + 1, cols=5)
    t_tf.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_tf)
    for c_idx, h in enumerate(tf_hdr):
        cell = t_tf.rows[0].cells[c_idx]
        cell.text = h
        set_cell_background(cell, "002855")
        set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
        p = cell.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(7.5)

    for r_idx, tfrow in enumerate(note.task_force_table):
        row = t_tf.rows[r_idx + 1]
        vals = [tfrow.sl_no, tfrow.material_code, tfrow.item_name, tfrow.abp_req_plus_safety, tfrow.recommended_qty]
        for c_idx, v in enumerate(vals):
            cell = row.cells[c_idx]
            cell.text = v
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            if r_idx % 2 == 1: set_cell_background(cell, "F8FAFC")
            row.cells[c_idx].paragraphs[0].runs[0].font.size = Pt(8)

    # Section 5: Cost Estimates Table
    p_sec5 = doc.add_paragraph()
    p_sec5.paragraph_format.space_before = Pt(10)
    p_sec5.paragraph_format.space_after = Pt(4)
    r5 = p_sec5.add_run("5. Basis of Estimated Value:\n")
    r5.bold = True
    r5.font.size = Pt(9)
    r5.font.color.rgb = navy
    r5_sub = p_sec5.add_run(note.cost_estimate_intro)
    r5_sub.font.size = Pt(8)

    cost_hdr = ["Sl. No.", "Description", "Unit", "Value"]
    t_cost = doc.add_table(rows=len(note.cost_estimate_table) + 1, cols=4)
    t_cost.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_cost)
    for c_idx, h in enumerate(cost_hdr):
        cell = t_cost.rows[0].cells[c_idx]
        cell.text = h
        set_cell_background(cell, "002855")
        set_cell_margins(cell, top=50, bottom=50, left=60, right=60)
        p = cell.paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
        p.runs[0].font.size = Pt(7.5)

    for r_idx, crow in enumerate(note.cost_estimate_table):
        row = t_cost.rows[r_idx + 1]
        vals = [crow.sl_no, crow.description, crow.unit, crow.value]
        for c_idx, v in enumerate(vals):
            cell = row.cells[c_idx]
            cell.text = v
            set_cell_margins(cell, top=40, bottom=40, left=60, right=60)
            if r_idx % 2 == 1: set_cell_background(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.runs[0].font.size = Pt(8)
            if crow.sl_no == "5": p.runs[0].font.bold = True

    # Section 6: Mode of Tender Justification
    p_sec6 = doc.add_paragraph()
    p_sec6.paragraph_format.space_before = Pt(10)
    p_sec6.paragraph_format.space_after = Pt(4)
    r6 = p_sec6.add_run("6. Justification for procurement of MS-Shredded Scrap through EPS (M-junction) on OTE basis:\n")
    r6.bold = True
    r6.font.size = Pt(9)
    r6.font.color.rgb = navy
    r6_sub = p_sec6.add_run(note.mode_tender_justification)
    r6_sub.font.size = Pt(8)

    # Proposals
    p_prop = doc.add_paragraph()
    p_prop.paragraph_format.space_before = Pt(10)
    p_prop.paragraph_format.space_after = Pt(4)
    r_p = p_prop.add_run("In view of the above, the following are proposed:")
    r_p.bold = True
    r_p.font.size = Pt(9.5)
    r_p.font.color.rgb = navy

    for prop in note.proposals:
        p_item = doc.add_paragraph(prop)
        p_item.paragraph_format.left_indent = Inches(0.2)
        p_item.paragraph_format.space_after = Pt(2)
        p_item.runs[0].font.size = Pt(8)

    # Approval Sought for
    p_app = doc.add_paragraph()
    p_app.paragraph_format.space_before = Pt(10)
    p_app.paragraph_format.space_after = Pt(3)
    r_app = p_app.add_run("Approval Sought for:\n")
    r_app.bold = True
    r_app.font.size = Pt(9)
    r_app.font.color.rgb = navy
    r_app_txt = p_app.add_run(note.approval_sought)
    r_app_txt.font.size = Pt(8)

    # DOP Hierarchy
    p_dop = doc.add_paragraph()
    p_dop.paragraph_format.space_before = Pt(10)
    p_dop.paragraph_format.space_after = Pt(3)
    r_dop = p_dop.add_run("DOP / Manual / Circular Ref & Approver:\n")
    r_dop.bold = True
    r_dop.font.size = Pt(9)
    r_dop.font.color.rgb = navy
    r_dop_sub = p_dop.add_run("As per Delegation of Power (DOP) and Office Order forwarded on behalf of GM I/c (SMS).\n")
    r_dop_sub.font.size = Pt(8)
    r_dop_auth = p_dop.add_run(note.dop_hierarchy)
    r_dop_auth.bold = True
    r_dop_auth.font.size = Pt(8)
    r_dop_auth.font.color.rgb = navy

    # Boilerplate Disclaimers
    p_disc = doc.add_paragraph()
    p_disc.paragraph_format.space_before = Pt(14)
    p_disc.paragraph_format.space_after = Pt(3)
    r_disc_title = p_disc.add_run("Standard Disclaimers and Usage Notes [TEMPLATE SECTION]\n")
    r_disc_title.bold = True
    r_disc_title.font.size = Pt(8)
    r_disc_title.font.color.rgb = RGBColor(100, 116, 139)
    r_disc_body = p_disc.add_run("This note is an automated reproduction of the uploaded document in a standard enquiry-proposal format and is provided for convenience only. It does not replace, amend or override the original document, its corrigenda, annexures or any instruction issued by the competent authority. Every figure, date, clause reference and approval must be verified against the original before being acted upon. Where the source was silent, ambiguous or illegible, the corresponding field is marked 'Not specified in the document' rather than estimated.\n\n— End of Report —")
    r_disc_body.font.size = Pt(7.5)
    r_disc_body.font.italic = True
    r_disc_body.font.color.rgb = RGBColor(100, 116, 139)

    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output